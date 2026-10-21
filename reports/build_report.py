"""Build FaceMatch_Analysis.xlsx from cached embeddings (work8k/) + earlier results (work/)."""
import glob, json, os

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.image import Image as XlImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TH = 0.40  # operating threshold used for the decision column
OUT = "FaceMatch_Analysis.xlsx"

# ---------- data ----------
sp = pd.read_csv("work8k/split.csv")
tr = sp[sp.split == "train"].reset_index(drop=True)
te = sp[sp.split == "test"].reset_index(drop=True)
un = sp[sp.split == "unknown"].reset_index(drop=True)
emb = {np.load(f)["out"].shape[0]: np.load(f)["out"] for f in glob.glob("work8k/emb_*.npz")}
G, P, U = emb[len(tr)], emb[len(te)], emb[len(un)]
short = lambda p: f"{os.path.basename(os.path.dirname(p))}/{os.path.basename(p)}"
g_name, g_person = tr.path.map(short).values, tr.person.values
true_idx = {p: i for i, p in enumerate(g_person)}


def topk(Q, k=5, truth=None):
    """Exact top-k per query (chunked). Returns idx[k], sims[k], and (if truth) true-person sim and rank."""
    I, S5, own, rank = [], [], [], []
    for s in range(0, len(Q), 2000):
        S = Q[s:s + 2000] @ G.T
        part = np.argpartition(-S, k, axis=1)[:, :k]
        order = np.take_along_axis(S, part, 1).argsort(1)[:, ::-1]
        idx = np.take_along_axis(part, order, 1)
        I.append(idx); S5.append(np.take_along_axis(S, idx, 1))
        if truth is not None:
            ti = np.array([true_idx[p] for p in truth[s:s + 2000]])
            o = S[np.arange(len(S)), ti]
            own.append(o); rank.append((S > o[:, None]).sum(1) + 1)
    r = [np.vstack(I), np.vstack(S5)]
    return r + ([np.concatenate(own), np.concatenate(rank)] if truth is not None else [])


tI, tS, own, rank = topk(P, truth=te.person.values)
uI, uS = topk(U)
right = g_person[tI[:, 0]] == te.person.values
top1s = tS[:, 0]
n, nu = len(te), len(un)

# ---------- styles ----------
HDR = PatternFill("solid", fgColor="1F3864"); HF = Font(bold=True, color="FFFFFF")
SUB = PatternFill("solid", fgColor="D9E2F3"); GOOD = PatternFill("solid", fgColor="C6EFCE")
BAD = PatternFill("solid", fgColor="FFC7CE"); WARN = PatternFill("solid", fgColor="FFEB9C")
PCT = "0.0%"; wrap = Alignment(wrap_text=True, vertical="top")


def header(ws, row, cols, start=1):
    for j, c in enumerate(cols):
        x = ws.cell(row, start + j, c); x.fill, x.font = HDR, HF
        x.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = Font(bold=True, size=16, color="1F3864")
    if sub:
        ws["A2"] = sub; ws["A2"].font = Font(italic=True, color="595959")


wb = Workbook()

# ---------- thresholds ----------
ths = [round(x, 2) for x in np.arange(0.20, 0.701, 0.05)]
o500 = pd.read_csv("work/open_set.csv").set_index("threshold")
rows_th = []
for t in ths:
    acc_ok = (right & (top1s >= t)).sum(); acc_bad = (~right & (top1s >= t)).sum()
    acc_unk = (uS[:, 0] >= t).sum()
    rows_th.append([t, acc_ok / n, acc_bad / n, ((right) & (top1s < t)).sum() / n, acc_unk / nu,
                    acc_ok / max(acc_ok + acc_bad, 1), acc_ok / max(acc_ok + acc_bad + acc_unk, 1),
                    float(o500.loc[t, "DIR (known found)"]), float(o500.loc[t, "FPIR (unknown accepted)"])])

# ---------- 1. Summary ----------
ws = wb.active; ws.title = "Summary"
title(ws, "Face Match (1:N) - Evaluation Summary", "antelopev2 (SCRFD + glintr100, 512-d) + Qdrant (Colima) | VGGFace2 112x112 | exact full-scan search")
i = ths.index(TH); r = rows_th[i]
top = lambda k: float((rank <= k).mean())  # rank = position of true person among all 8k
ws["A4"] = "Test setup"; ws["A4"].font = Font(bold=True, size=12)
setup = [("Enrolled people (gallery)", f"{len(tr):,} - 1 photo each"), ("Test photos (same people)", f"{n:,} (3 per person)"),
         ("Unknown people (never enrolled)", f"{nu:,} photos / 500 people"), ("Search", "Exact full scan (cosine); HNSW compared separately"),
         ("Operating threshold", f"cosine similarity >= {TH:.2f}")]
for k, (a, b) in enumerate(setup, 5):
    ws.cell(k, 1, a); ws.cell(k, 2, b)
ws["A11"] = "Headline results"; ws["A11"].font = Font(bold=True, size=12)
header(ws, 12, ["Metric", "Value", "What it means"])
kp = [("Top-1 accuracy", float(right.mean()), "Right person ranked first"),
      ("Top-5 accuracy", top(5), "Right person among first 5 candidates"),
      ("Top-10 accuracy", top(10), "Right person among first 10 candidates"),
      (f"Found automatically @ {TH}", r[1], "Enrolled people matched correctly and accepted"),
      (f"Precision of accepted matches @ {TH}", r[5], "Of matches the system accepts, share that are correct (enrolled probes only)"),
      (f"Precision incl. unknown people @ {TH}", r[6], "Same, counting unknown people wrongly accepted (depends on traffic mix)"),
      (f"Confident wrong match @ {TH}", r[2], "Wrong person named with score >= threshold - the risky error"),
      (f"Unknown person wrongly accepted @ {TH}", r[4], "Never-enrolled probes that scored >= threshold"),
      ("Embedding latency (1 face, CPU) p50", "102 ms", "glintr100 on 10-core Apple Silicon (p95 144 ms)"),
      ("Search latency (8k vectors) p50", "~3 ms", "Qdrant top-20, exact scan, incl. local HTTP round trip")]
for k, (a, b, c) in enumerate(kp, 13):
    ws.cell(k, 1, a); x = ws.cell(k, 2, b); ws.cell(k, 3, c)
    if isinstance(b, float): x.number_format = PCT
    x.font = Font(bold=True)
ws["A24"] = "Key findings"; ws["A24"].font = Font(bold=True, size=12)
finds = [f"1. The model finds the right person first {right.mean():.1%} of the time with 1 enrolled photo and {len(tr):,} people (95.3% with 500 people).",
         f"2. At threshold {TH}, {r[5]:.1%} of accepted matches are correct, but {1 - r[1]:.1%} of enrolled people are not matched automatically (send to review).",
         f"3. Most misses are 'not found' ({(~right & (top1s < TH)).mean():.1%}), not wrong names; only {r[2]:.2%} are confident wrong matches.",
         "4. Top-5/Top-10 add only ~1 point over Top-1: the misses are mostly hard or mislabeled photos, not near misses.",
         "5. HNSW matches exact search once hnsw_ef >= 128 (99.9% same top-1); at 8k-100k vectors it is not needed for speed.",
         "6. Exact scan latency grows linearly (~0.066 ms per 1k vectors); HNSW stays ~flat (3.3 -> 3.6 ms from 8k to 100k).",
         "7. Embedding (~100 ms) dominates per-request latency, not vector search."]
for k, t in enumerate(finds, 25):
    ws.cell(k, 1, t)
ws["A33"] = "Recommendations"; ws["A33"].font = Font(bold=True, size=12)
recs = ["- Use 0.45-0.50 as 'likely same person' for dedupe (wrong-accept <1%); 0.35-0.45 -> manual review.",
        "- Enrol several photos per person (or average embeddings) to recover false negatives.",
        "- Move to HNSW only when the collection is large (hundreds of thousands+); use hnsw_ef >= 128 and re-measure recall at real size.",
        "- Validate on real Cashfree data (ID photo vs selfie) before fixing thresholds: VGGFace2 is the easy case.",
        "- Confirm licensing: antelopev2 weights and VGGFace2 are research / non-commercial."]
for k, t in enumerate(recs, 34):
    ws.cell(k, 1, t)
widths(ws, [46, 26, 80])

# ---------- 2. Accuracy vs scale ----------
ws = wb.create_sheet("Accuracy vs Scale")
title(ws, "Closed-set accuracy vs gallery size", "1 enrolled photo per person, exact search")
s500 = json.load(open("work/summary.json"))
header(ws, 4, ["Metric", "500 people (5,000 probes)", f"{len(tr):,} people ({n:,} probes)", "Change (pts)"])
for k, (m, a, b) in enumerate([("Top-1", s500["top1"], float(right.mean())), ("Top-5", s500["top5"], top(5)),
                               ("Top-10", s500["top10"], top(10))], 5):
    ws.cell(k, 1, m); ws.cell(k, 2, a).number_format = PCT; ws.cell(k, 3, b).number_format = PCT
    ws.cell(k, 4, (b - a) * 100).number_format = "0.0"
ch = BarChart(); ch.type = "col"; ch.title = "Top-K accuracy by gallery size"; ch.y_axis.scaling.min = 0.9; ch.y_axis.scaling.max = 1.0
ch.add_data(Reference(ws, min_col=2, max_col=3, min_row=4, max_row=7), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=7)); ch.height, ch.width = 8, 16
ws.add_chart(ch, "A10"); widths(ws, [18, 26, 26, 14])

# ---------- 3. Threshold analysis ----------
ws = wb.create_sheet("Threshold Analysis")
title(ws, f"Threshold trade-off ({len(tr):,} people)", "Pick the row that matches your risk appetite. Similarity = cosine of face embeddings (not a probability).")
header(ws, 4, ["Threshold (cosine)", "Found & accepted", "Confident WRONG match", "Correct but rejected", "Unknown wrongly accepted",
               "Precision (known probes)", "Precision (incl. unknown)", "[500 ppl] Found", "[500 ppl] Unknown accepted"])
for k, rw in enumerate(rows_th, 5):
    for j, v in enumerate(rw, 1):
        c = ws.cell(k, j, v); c.number_format = "0.00" if j == 1 else PCT
    if rw[0] == TH:
        for j in range(1, 10): ws.cell(k, j).fill = WARN
ws.cell(5 + len(rows_th) + 1, 1, f"Highlighted row = operating point used in the Photo-wise sheet ({TH}).")
lc = LineChart(); lc.title = "Found vs unknown-accepted by threshold"; lc.height, lc.width = 9, 18
lc.add_data(Reference(ws, min_col=2, min_row=4, max_row=4 + len(ths)), titles_from_data=True)
lc.add_data(Reference(ws, min_col=5, min_row=4, max_row=4 + len(ths)), titles_from_data=True)
lc.set_categories(Reference(ws, min_col=1, min_row=5, max_row=4 + len(ths))); ws.add_chart(lc, "A20")
widths(ws, [14, 16, 18, 16, 18, 18, 18, 16, 18]); ws.row_dimensions[4].height = 45

# ---------- 4. HNSW vs exact ----------
ws = wb.create_sheet("HNSW vs Exact")
title(ws, "HNSW vs exact scan (8,000 vectors, 3,000 probes)", "HNSW m=16, ef_construct=100. 'ef' = candidate list size at search time (hnsw_ef).")
header(ws, 4, ["Search", "Top-1 accuracy", "Same top-1 as exact", "p50 latency (ms)"])
hn = [("Exact scan", .9353, 1.0, 3.0), ("HNSW ef=16", .9150, .9613, 2.5), ("HNSW ef=64", .9337, .9953, 3.0),
      ("HNSW ef=128", .9350, .9987, 3.3), ("HNSW ef=256", .9353, 1.0, 4.1)]
for k, rw in enumerate(hn, 5):
    for j, v in enumerate(rw, 1):
        c = ws.cell(k, j, v); c.number_format = {2: "0.00%", 3: "0.00%", 4: "0.0"}.get(j, "General")
for k, t in enumerate(["Reading: low ef is fast but can stop at a near-best face (-2 pts at ef=16). ef>=128 is practically identical to exact.",
                       "HNSW only changes which stored vectors are compared; it cannot beat the exact result.",
                       "At 8k vectors HNSW gives no speed gain; the benefit appears at much larger sizes (see Latency & Scale).",
                       "Note: exact top-1 here (93.5%) is on a 3,000-probe subset; all 24,000 probes give 93.0%.",
                       "Recall at larger sizes was NOT measured - re-test at the real collection size."], 11):
    ws.cell(k, 1, t)
widths(ws, [22, 18, 22, 18])

# ---------- 5. Latency & scale ----------
ws = wb.create_sheet("Latency & Scale")
title(ws, "Latency", "Measured single-query, p50; 10-core Apple Silicon Mac, Qdrant in Colima via localhost; top-20 results")
header(ws, 4, ["Collection size", "Exact scan (ms)", "HNSW ef=128 (ms)"])
for k, rw in enumerate([(8000, 3.0, 3.3), (25000, 4.1, 3.4), (100000, 9.1, 3.6)], 5):
    for j, v in enumerate(rw, 1): ws.cell(k, j, v)
ws.cell(8, 1, 1000000); ws.cell(8, 2, 70).fill = WARN; ws.cell(8, 3, "low single digits").fill = WARN
ws.cell(9, 1, "Row 1,000,000 is an extrapolation, NOT measured (exact ~2.5 ms fixed + 0.066 ms per 1k vectors).")
ws.cell(10, 1, "Padding vectors for 25k/100k were random unit vectors (latency only; recall not meaningful).")
header(ws, 12, ["Step", "p50 (ms)", "p95 (ms)"])
for k, rw in enumerate([("Embed 1 face (glintr100, CPU)", 102, 144), ("Search 8k, exact", 2.9, 3.8), ("Search 8k, default", 2.8, 3.4)], 13):
    for j, v in enumerate(rw, 1): ws.cell(k, j, v)
ws.cell(17, 1, "Face detection + alignment (SCRFD) NOT timed: VGGFace2 images are pre-aligned. Expect it to add on top of embedding.")
lc = LineChart(); lc.title = "Search latency vs collection size (p50 ms)"; lc.height, lc.width = 8, 16
lc.add_data(Reference(ws, min_col=2, max_col=3, min_row=4, max_row=7), titles_from_data=True)
lc.set_categories(Reference(ws, min_col=1, min_row=5, max_row=7)); ws.add_chart(lc, "E4")
widths(ws, [34, 18, 18])

# ---------- 6. Failure analysis ----------
ws = wb.create_sheet("Failure Analysis")
title(ws, f"Where the errors come from (threshold {TH}, {n:,} probes)")
header(ws, 3, ["Outcome", "Count", "Share", "Risk"])
cats = [("Correct, accepted", right & (top1s >= TH), "OK"), ("Correct but rejected (score below threshold)", right & (top1s < TH), "Low - goes to review"),
        ("Wrong top-1, rejected ('not found')", ~right & (top1s < TH), "Low - goes to review"),
        ("Wrong top-1, ACCEPTED (confident wrong match)", ~right & (top1s >= TH), "HIGH")]
for k, (a, m, rk) in enumerate(cats, 4):
    ws.cell(k, 1, a); ws.cell(k, 2, int(m.sum())); ws.cell(k, 3, m.mean()).number_format = "0.00%"; ws.cell(k, 4, rk)
    ws.cell(k, 4).fill = BAD if rk == "HIGH" else (GOOD if rk == "OK" else WARN)
ws.cell(9, 1, "Similarity of the TRUE person among top-1 misses").font = Font(bold=True)
header(ws, 10, ["Range", "Count", "Share of misses"])
ms = own[~right]
for k, (a, m) in enumerate([("< 0.20 (model sees different faces)", ms < .2), ("0.20 - 0.40", (ms >= .2) & (ms < .4)), (">= 0.40", ms >= .4)], 11):
    ws.cell(k, 1, a); ws.cell(k, 2, int(m.sum())); ws.cell(k, 3, m.mean()).number_format = PCT
ws.cell(15, 1, "Rank of the true person (all probes)").font = Font(bold=True)
header(ws, 16, ["Rank", "Count", "Share of probes"])
for k, (a, m) in enumerate([("1", rank == 1), ("2", rank == 2), ("3-5", (rank >= 3) & (rank <= 5)), ("6-10", (rank >= 6) & (rank <= 10)),
                            ("11-100", (rank >= 11) & (rank <= 100)), ("> 100", rank > 100)], 17):
    ws.cell(k, 1, a); ws.cell(k, 2, int(m.sum())); ws.cell(k, 3, m.mean()).number_format = "0.00%"
ws.cell(24, 1, "Interpretation: misses are mostly not near misses (true person's median similarity ~0.13) - likely hard/mislabeled photos (not verified by eye).")
widths(ws, [52, 12, 16, 24])

# ---------- 7. Photo-wise ----------
ws = wb.create_sheet("Photo-wise (24k)")
cols = ["Probe image", "True person"]
for k in range(1, 6): cols += [f"Top{k} person", f"Top{k} image", f"Top{k} similarity"]
cols += ["True person rank", "True person similarity", "Top-1 correct", "In top-5", f"Decision @ {TH}", "Confidence band"]
header(ws, 1, cols)
probe_name = te.path.map(short).values; truth = te.person.values
for j in range(n):
    row = [probe_name[j], truth[j]]
    for k in range(5):
        row += [g_person[tI[j, k]], g_name[tI[j, k]], float(tS[j, k])]
    ok1 = bool(right[j]); s1 = top1s[j]
    dec = ("Correct match (accepted)" if ok1 and s1 >= TH else "WRONG match (accepted)" if not ok1 and s1 >= TH
           else "Correct but rejected" if ok1 else "Not found (rejected)")
    band = "High (>=0.50)" if s1 >= .5 else "Medium (0.40-0.50)" if s1 >= TH else "Low (<0.40)"
    row += [int(rank[j]), float(own[j]), "Y" if ok1 else "N", "Y" if rank[j] <= 5 else "N", dec, band]
    ws.append(row)
for col in (5, 8, 11, 14, 17, 19):
    for c in ws.iter_cols(min_col=col, max_col=col, min_row=2):
        for x in c: x.number_format = PCT
dcol = get_column_letter(len(cols) - 1)
for txt, fill in (("WRONG", BAD), ("Not found", WARN), ("Correct match", GOOD)):
    ws.conditional_formatting.add(f"{dcol}2:{dcol}{n + 1}", FormulaRule(formula=[f'ISNUMBER(SEARCH("{txt}",{dcol}2))'], fill=fill))
ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
widths(ws, [22, 12] + [12, 22, 11] * 5 + [10, 11, 9, 9, 24, 18]); ws.row_dimensions[1].height = 32

# ---------- 8. Unknown people ----------
ws = wb.create_sheet("Unknown people (1.5k)")
header(ws, 1, ["Probe image", "Person (NOT enrolled)"] + [x for k in range(1, 6) for x in (f"Top{k} person", f"Top{k} image", f"Top{k} similarity")] + [f"Wrongly accepted @ {TH}"])
un_name = un.path.map(short).values
for j in range(nu):
    row = [un_name[j], un.person[j]]
    for k in range(5): row += [g_person[uI[j, k]], g_name[uI[j, k]], float(uS[j, k])]
    row.append("YES" if uS[j, 0] >= TH else "no"); ws.append(row)
for col in (5, 8, 11, 14, 17):
    for c in ws.iter_cols(min_col=col, max_col=col, min_row=2):
        for x in c: x.number_format = PCT
ws.conditional_formatting.add(f"R2:R{nu + 1}", FormulaRule(formula=['R2="YES"'], fill=BAD))
ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions; widths(ws, [22, 14] + [12, 22, 11] * 5 + [18])

# ---------- 9. Confident wrong (with thumbnails) ----------
ws = wb.create_sheet("Confident wrong (images)")
bad = np.where(~right & (top1s >= TH))[0]; bad = bad[np.argsort(-top1s[bad])]
title(ws, f"Confident wrong matches @ {TH}: {len(bad)} cases (highest similarity first)", "Check by eye: some may be mislabeled ground truth or genuinely look-alike faces.")
header(ws, 3, ["Probe photo", "Matched (enrolled) photo", "Probe image", "True person", "Matched person", "Similarity", "True person's similarity", "True person rank"])
for k, j in enumerate(bad, 4):
    ws.row_dimensions[k].height = 62
    for col, path in ((1, te.path[j]), (2, tr.path[tI[j, 0]])):
        im = XlImage(path); im.width = im.height = 80; ws.add_image(im, f"{get_column_letter(col)}{k}")
    for col, v in enumerate([probe_name[j], truth[j], g_person[tI[j, 0]], float(top1s[j]), float(own[j]), int(rank[j])], 3):
        ws.cell(k, col, v).alignment = Alignment(vertical="center")
    ws.cell(k, 6).number_format = PCT; ws.cell(k, 7).number_format = PCT
ws.freeze_panes = "A4"; widths(ws, [14, 24, 22, 14, 14, 12, 14, 12])

# ---------- 10. Method ----------
ws = wb.create_sheet("Method & Caveats")
title(ws, "Method and caveats")
lines = ["PIPELINE: image -> (VGGFace2 images are pre-aligned 112x112, detection skipped) -> glintr100 (ArcFace R100) 512-d L2-normalised embedding -> Qdrant cosine search.",
         "In production on raw photos: SCRFD detection + 5-point alignment (norm_crop) runs first via insightface FaceAnalysis.",
         "GALLERY: 8,000 identities, 1 photo each. TEST: 3 other photos per identity (24,000). UNKNOWN: 500 never-enrolled identities (1,500 photos).",
         "SEARCH: exact full scan (SearchParams exact=True). No HNSW graph was used for the headline results.",
         "SIMILARITY: cosine similarity shown as %. It is NOT the probability the match is correct.",
         "PRECISION: computed on accepted matches at the chosen threshold; 'incl. unknown' depends on the share of unknown people in real traffic.",
         "LIMITS: VGGFace2 is cleaner than real KYC photos; single random gallery/seed; failure causes (pose/blur/mislabel) not verified by eye.",
         "HNSW recall tested only at 8,000 vectors on 3,000 probes. 25k/100k latency used random padding vectors; 1M is extrapolated.",
         "LICENCE: antelopev2 weights and VGGFace2 are for non-commercial research use - confirm before production use.",
         "ENVIRONMENT: 10-core Apple Silicon Mac, ONNX Runtime CPU, Qdrant in Colima; latency includes local HTTP round trip."]
for k, t in enumerate(lines, 3):
    ws.cell(k, 1, t).alignment = wrap
widths(ws, [150])

wb.save(OUT)
print("saved", OUT, os.path.getsize(OUT) // 1024, "KB")
print("confident wrong:", len(bad), "| top1", right.mean())
